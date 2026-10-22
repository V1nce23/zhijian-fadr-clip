"""Direct PDF and editable Word exports from the same immutable report snapshot."""

import base64
import io
import re
from functools import lru_cache
from pathlib import Path
from xml.sax.saxutils import escape

from app.services.report_export import (
    VERDICTS, STATUSES, FRAME_STATUSES, _date, _number, _score, _image_data,
)
from app.services.evidence_gallery import gallery_prefix

FONT_PATH = Path(__file__).resolve().parents[1] / "assets" / "fonts" / "NotoSansSC-Regular.ttf"
FONT_NAME = "ZhijianSans"


def text(value):
    value = "未记录" if value is None else str(value)
    value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", "", value)
    return value.replace("\u2011", "-").replace("\u2013", "-").replace("\u2014", "-")


def policy_text(data):
    policy = data["record"].get("threshold_policy")
    if not policy:
        return "历史模式 沿用原报告结论和已记录阈值"
    unit = "视频均值" if policy["score_unit"] == "video_mean" else "图片分数"
    return f'{policy["label"]}模式 {unit} > {_number(policy["threshold"])}'


def threshold_text(extra):
    thresholds = extra["thresholds"]
    low, high = thresholds.get("authentic_max"), thresholds.get("manipulated_min")
    if low is None or high is None:
        return "单帧阈值未记录 不推算缺失结论"
    if thresholds.get("operator") == ">":
        return f"图片或单帧分数 > {_number(high)} 判为疑似伪造 否则为疑似真实"
    return f"历史规则 分数 >= {_number(high)} 为疑似伪造 <= {_number(low)} 为疑似真实 其间为不确定"


def metadata(data):
    record, upload = data["record"], data["record"]["upload"]
    devices = {"cpu": "CPU", "cuda": "GPU"}
    return [
        ("记录 ID", record["id"]), ("文件 SHA-256", upload["sha256"]),
        ("类型和大小", f'{ {"video": "视频", "image": "图片", "audio": "历史音频"}.get(upload["media_type"], upload["media_type"])} / {upload["size_bytes"]} 字节'),
        ("提交时间", _date(record["created_at"])), ("完成时间", _date(record["completed_at"])),
        ("所选和实际设备", f'{devices.get(record["requested_device"], "未记录")} / {devices.get(record["executed_device"], "未记录")}'),
    ]


def frame_rows(extra, technical=True):
    rows = [["序号", "时间", "状态", "分数", "结论", "L6", "L12", "L18", "L24", "路由来源"]]
    for frame in extra["frames"]:
        route = frame.get("routing")
        rows.append([str(frame["frame_number"]), _number(frame["timestamp_seconds"]), FRAME_STATUSES.get(frame["status"], frame["status"]),
            _score(frame["score"]), VERDICTS.get(frame["verdict"], "-"),
            *([f"{value * 100:.2f}%" for value in route["weights"]] if route else ["-"] * 4),
            "补充生成" if route and route["origin"] == "supplementary_replay" else "原检测" if route else "未记录"])
    return [[text(cell) for cell in (row if technical else row[:5])] for row in rows]


def examples(data, run, extra, storage):
    """Use at most three cached 256px inputs; never queue or re-run a detector."""
    if data["record"]["upload"]["status"] == "deleted":
        return []
    from PIL import Image
    from app.services.attention_gallery import prefix, read_manifest

    assets = {}
    try:
        assets = {item["frame_index"]: item for item in read_manifest(storage, data["record"]["id"], run["id"])["frames"]}
    except Exception:
        pass
    result = []
    for frame in sorted((item for item in extra["frames"] if item["score"] is not None), key=lambda item: item["score"], reverse=True)[:3]:
        try:
            face_uri = _image_data(storage, f'{gallery_prefix(data["record"]["id"], run["id"])}frame-{frame["frame_index"]}.png')
            if not face_uri:
                continue
            face_bytes = base64.b64decode(face_uri.split(",", 1)[1])
            face = Image.open(io.BytesIO(face_bytes))
            if face.size != (256, 256):
                continue
            overlay = None
            asset = assets.get(frame["frame_index"])
            heat_uri = _image_data(storage, f'{prefix(data["record"]["id"], run["id"])}frame-{frame["frame_index"]}.png') if asset else None
            if heat_uri:
                heat = Image.open(io.BytesIO(base64.b64decode(heat_uri.split(",", 1)[1])))
                if heat.size == (256, 256):
                    heat = heat.convert("RGBA")
                    heat.putalpha(heat.getchannel("A").point(lambda alpha: round(alpha * 0.55)))
                    composed = Image.alpha_composite(face.convert("RGBA"), heat).convert("RGB")
                    stream = io.BytesIO()
                    composed.save(stream, format="PNG")
                    overlay = stream.getvalue()
            result.append({"frame": frame, "face": face_bytes, "overlay": overlay,
                "attention_origin": "补充生成" if asset and asset.get("origin") == "supplementary_replay" else "原检测"})
        except Exception:
            continue
    return result


def routing_caption(routing):
    if not routing:
        return "本记录未保存可用路由权重"
    values = "  ".join(f'L{layer} {weight * 100:.2f}%' for layer, weight in zip(routing["layers"], routing["weights"]))
    origin = "补充生成" if routing.get("origin") == "supplementary_replay" else "原检测"
    return f"{values}  来源 {origin}"


def ending(data):
    return [data["interpretation"], f'导出时间 {_date(data["exported_at"])} 此文件为导出时快照 补充数据不改写原检测结果',
            "原报告数据校验值 SHA-256", data["record_sha256"]]


@lru_cache(maxsize=1)
def register_pdf_font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    pdfmetrics.registerFont(TTFont(FONT_NAME, str(FONT_PATH)))
    pdfmetrics.registerFontFamily(FONT_NAME, normal=FONT_NAME, bold=FONT_NAME, italic=FONT_NAME, boldItalic=FONT_NAME)


def pdf_report(data, storage):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, KeepTogether, PageBreak, Flowable

    register_pdf_font()
    technical = data.get("technical_details_included", True)
    output = io.BytesIO()
    width = letter[0] - 96
    styles = {
        "body": ParagraphStyle("body", fontName=FONT_NAME, fontSize=10, leading=13, spaceAfter=4, wordWrap="CJK"),
        "small": ParagraphStyle("small", fontName=FONT_NAME, fontSize=8.3, leading=10, spaceAfter=3, wordWrap="CJK"),
        "cell": ParagraphStyle("cell", fontName=FONT_NAME, fontSize=8.1, leading=10.5, alignment=TA_CENTER, wordWrap="CJK"),
        "title": ParagraphStyle("title", fontName=FONT_NAME, fontSize=19, leading=25, spaceAfter=8, textColor=colors.black, wordWrap="CJK"),
        "heading": ParagraphStyle("heading", fontName=FONT_NAME, fontSize=12, leading=17, spaceBefore=10, spaceAfter=7, keepWithNext=True, wordWrap="CJK"),
    }

    def p(value, style="body"):
        return Paragraph(escape(text(value)).replace("\n", "<br/>"), styles[style])

    class RoutingChart(Flowable):
        def __init__(self, routing):
            super().__init__()
            self.routing = routing
            self.width, self.height = width, 67

        def draw(self):
            canvas = self.canv
            for index, (layer, weight) in enumerate(zip(self.routing["layers"], self.routing["weights"])):
                y = 53 - index * 16
                canvas.setFont(FONT_NAME, 9)
                canvas.setFillColor(colors.black)
                canvas.drawString(0, y, f"L{layer}")
                canvas.setFillColor(colors.HexColor("#e8edf4"))
                canvas.roundRect(38, y, width - 114, 8, 3, fill=1, stroke=0)
                canvas.setFillColor(colors.HexColor("#2879c9"))
                if weight > 0:
                    canvas.roundRect(38, y, (width - 114) * weight, 8, min(3, (width - 114) * weight / 2), fill=1, stroke=0)
                canvas.setFillColor(colors.black)
                canvas.drawRightString(width, y, f"{weight * 100:.2f}%")

    record = data["record"]
    story = [p("智鉴伪造人脸分析报告", "title"), p(record["upload"]["original_name"]),
        p(f'总体结论 {VERDICTS.get(record["verdict"], "未生成结论")}  伪造倾向分数 {_score(record["score"])}'),
        p(record["summary"] or "本次任务未生成检测结论"),
        p(f'状态 {STATUSES.get(record["status"], record["status"])}  判断模式 {policy_text(data)}', "small")]
    if record["score"] is not None:
        story.append(p(f'原始检测分数 {_number(record["score"])}', "small"))
    meta = Table([[p(label, "small"), p(value, "small")] for label, value in metadata(data)], colWidths=[84, width - 84], hAlign="LEFT")
    meta.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
    story.extend([meta, Spacer(1, 6)])

    for run, extra in zip(record["detector_runs"], data["supplementary"]):
        model = "FADR-CLIP" if run["detector_id"] == "fadr-clip" else run["detector_id"]
        story.extend([p(f"{model} 模型结果", "heading"), p(f'{VERDICTS.get(run["verdict"], "未生成结论")} {_score(run["score"])}  {STATUSES.get(run["status"], run["status"])}', "small")])
        if technical:
            story.extend([p(f'模型版本 {run["model_version"]}  预处理版本 {run["preprocessing_version"]}', "small"),
                p("有效人脸帧平均路由权重" if record["upload"]["media_type"] == "video" else "图片路由权重", "heading")])
            story.append(RoutingChart(extra["routing_mean"]) if extra["routing_mean"] else p("本记录未保存可用路由权重", "small"))
        story.append(p(threshold_text(extra), "small"))
        if run.get("error"):
            story.append(p(f'模型记录 {run["error"]}', "small"))
        samples = examples(data, run, extra, storage)
        large = len(extra["frames"]) > 8
        if samples:
            if large and technical:
                story.append(PageBreak())
            story.append(p("有效人脸可视化样本", "heading"))
            story.append(p("按原检测分数从高到低选取 最多三张 仅使用已有缓存", "small"))
            for sample in samples:
                frame = sample["frame"]
                label = f'第 {frame["frame_number"]} 帧' if record["upload"]["media_type"] == "video" else "检测人脸"
                size = 92 if large else 72
                pictures = [Image(io.BytesIO(sample["face"]), size, size)]
                captions = [p("对齐后检测人脸", "small")]
                if sample["overlay"]:
                    pictures.append(Image(io.BytesIO(sample["overlay"]), size, size))
                    captions.append(p(f'路由空间注意力 {sample["attention_origin"]} 55% 叠加', "small"))
                image_table = Table([pictures, captions], colWidths=[width / len(pictures)] * len(pictures))
                image_table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("ALIGN", (0, 0), (-1, 0), "CENTER"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 2), ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
                story.append(KeepTogether([p(f'{label}  原检测分数 {_score(frame["score"])}', "small"), image_table]
                    + ([p(routing_caption(frame.get("routing")), "small")] if technical else []) + [Spacer(1, 6)]))
        else:
            story.append(p("未附加人脸图片 现有缓存缺失或原素材已删除 导出不重新检测", "small"))
        if large:
            story.append(PageBreak())
        story.extend([p("完整逐帧记录" if record["upload"]["media_type"] == "video" else "图片检测明细", "heading"),
            p("序号从 1 起计 时间单位为秒 分数按百分比显示 无人脸和未完成项不计为零", "small")])
        rows = frame_rows(extra, technical)
        if len(rows) > 1:
            table = Table([[p(cell, "cell") for cell in row] for row in rows], repeatRows=1,
                          colWidths=[32, 44, 68, 50, 62, 46, 46, 46, 46, 76] if technical else [50, 80, 150, 90, 146], hAlign="LEFT")
            table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d9d9d9")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8f0fa")), ("TEXTCOLOR", (0, 0), (-1, -1), colors.black),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3), ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
            story.append(table)
        else:
            story.append(p("此记录没有保存逐帧明细", "small"))
    story.append(Spacer(1, 6))
    story.extend(p(line, "small") for line in ending(data))

    def footer(canvas, document):
        canvas.setFont(FONT_NAME, 8)
        canvas.drawString(48, 24, "智鉴分析报告")
        canvas.drawRightString(letter[0] - 48, 24, f"第 {document.page} 页")

    document = SimpleDocTemplate(output, pagesize=letter, rightMargin=48, leftMargin=48, topMargin=38, bottomMargin=40,
        title="智鉴伪造人脸分析报告", author="智鉴", creator="智鉴报告导出")
    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def _routing_png(routing):
    from PIL import Image, ImageDraw, ImageFont
    image = Image.new("RGB", (1400, 250), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(str(FONT_PATH), 29)
    for index, (layer, weight) in enumerate(zip(routing["layers"], routing["weights"])):
        y = 8 + index * 60
        draw.text((8, y), f"L{layer}", font=font, fill="black")
        draw.rounded_rectangle((115, y + 10, 1120, y + 32), radius=10, fill="#e8edf4")
        if weight > 0:
            draw.rectangle((115, y + 10, 115 + 1005 * weight, y + 32), fill="#2879c9")
        draw.text((1195, y), f"{weight * 100:.2f}%", font=font, fill="black")
    stream = io.BytesIO()
    image.save(stream, format="PNG")
    return stream.getvalue()


def docx_report(data, storage):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    technical = data.get("technical_details_included", True)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(0.55), Inches(0.55)
    section.left_margin, section.right_margin = Inches(0.65), Inches(0.65)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Caption"):
        style = doc.styles[name]
        style.font.name = "Noto Sans SC"
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Noto Sans SC")
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.line_spacing = 1.0
        style.font.bold = name in ("Heading 1", "Heading 2")
        style.font.italic = False
        style.paragraph_format.keep_with_next = name in ("Title", "Heading 1", "Heading 2")
    # Templates may carry decorative borders or inherited caption emphasis.
    for style in doc.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    doc.styles["Normal"].font.size = Pt(9.5)
    doc.styles["Title"].font.size = Pt(20)
    doc.styles["Heading 1"].font.size = Pt(13)
    doc.styles["Heading 2"].font.size = Pt(11)
    doc.styles["Caption"].font.size = Pt(8)
    doc.styles["Caption"].paragraph_format.space_after = Pt(2)
    for name in ("Heading 1", "Heading 2"):
        doc.styles[name].paragraph_format.space_before = Pt(6)
        doc.styles[name].paragraph_format.space_after = Pt(4)
    doc.core_properties.title = "智鉴伪造人脸分析报告"
    doc.core_properties.author = "智鉴"
    doc.core_properties.last_modified_by = "智鉴"
    doc.core_properties.subject = "既有检测记录的报告快照"

    def paragraph(value, style=None):
        return doc.add_paragraph(text(value), style)

    def table(rows, widths, font_size=8.3):
        grid = doc.add_table(rows=1, cols=len(rows[0]))
        grid.autofit = False
        for column, width in zip(grid.columns, widths):
            column.width = Inches(width)
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            node = OxmlElement("w:" + edge)
            for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                node.set(qn("w:" + key), value)
            borders.append(node)
        grid._tbl.tblPr.append(borders)
        for index, row_data in enumerate(rows):
            row = grid.rows[0] if index == 0 else grid.add_row()
            no_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(no_split)
            if index == 0:
                row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for cell, value, width in zip(row.cells, row_data, widths):
                cell.width = Inches(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                props = cell._tc.get_or_add_tcPr()
                margins = OxmlElement("w:tcMar")
                for edge in ("top", "left", "bottom", "right"):
                    node = OxmlElement("w:" + edge)
                    node.set(qn("w:w"), "35" if edge in ("top", "bottom") else "50")
                    node.set(qn("w:type"), "dxa")
                    margins.append(node)
                props.append(margins)
                if index == 0:
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:fill"), "E8F0FA")
                    props.append(shading)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                run = p.add_run(text(value))
                run.font.size = Pt(font_size)
                run.font.color.rgb = RGBColor(0, 0, 0)
                run.bold = index == 0
        paragraph("").paragraph_format.space_after = Pt(1)
        return grid

    record = data["record"]
    paragraph("智鉴伪造人脸分析报告", "Title")
    paragraph(record["upload"]["original_name"])
    paragraph(f'总体结论 {VERDICTS.get(record["verdict"], "未生成结论")}  伪造倾向分数 {_score(record["score"])}')
    paragraph(record["summary"] or "本次任务未生成检测结论")
    paragraph(f'状态 {STATUSES.get(record["status"], record["status"])}  判断模式 {policy_text(data)}')
    if record["score"] is not None:
        paragraph(f'原始检测分数 {_number(record["score"])}', "Caption")
    for label, value in metadata(data):
        paragraph(f'{label}  {text(value)}', "Caption")
    for run, extra in zip(record["detector_runs"], data["supplementary"]):
        model = "FADR-CLIP" if run["detector_id"] == "fadr-clip" else run["detector_id"]
        paragraph(f"{model} 模型结果", "Heading 1")
        paragraph(f'{VERDICTS.get(run["verdict"], "未生成结论")} {_score(run["score"])}  {STATUSES.get(run["status"], run["status"])}')
        if technical:
            paragraph(f'模型版本 {run["model_version"]}  预处理版本 {run["preprocessing_version"]}', "Caption")
            paragraph("有效人脸帧平均路由权重" if record["upload"]["media_type"] == "video" else "图片路由权重", "Heading 2")
            if extra["routing_mean"]:
                p = paragraph("")
                p.add_run().add_picture(io.BytesIO(_routing_png(extra["routing_mean"])), width=Inches(6.3))
                paragraph(routing_caption(extra["routing_mean"]), "Caption")
            else:
                paragraph("本记录未保存可用路由权重")
        paragraph(threshold_text(extra), "Caption")
        if run.get("error"):
            paragraph(f'模型记录 {run["error"]}', "Caption")
        samples = examples(data, run, extra, storage)
        large = len(extra["frames"]) > 8
        if samples:
            if large and technical:
                doc.add_page_break()
            paragraph("有效人脸可视化样本", "Heading 1")
            paragraph("按原检测分数从高到低选取 最多三张 仅使用已有缓存", "Caption")
            for sample in samples:
                frame = sample["frame"]
                label = f'第 {frame["frame_number"]} 帧' if record["upload"]["media_type"] == "video" else "检测人脸"
                p = paragraph(f'{label}  原检测分数 {_score(frame["score"])}', "Heading 2")
                p.paragraph_format.keep_with_next = True
                p = paragraph("")
                p.paragraph_format.keep_with_next = True
                picture_width = Inches(1.2 if large else 1.0)
                p.add_run().add_picture(io.BytesIO(sample["face"]), width=picture_width)
                if sample["overlay"]:
                    p.add_run("    ")
                    p.add_run().add_picture(io.BytesIO(sample["overlay"]), width=picture_width)
                label = f'左为对齐人脸 右为路由空间注意力 {sample["attention_origin"]} 55% 叠加' if sample["overlay"] else '对齐后检测人脸'
                p = paragraph(label, "Caption")
                p.paragraph_format.keep_with_next = technical
                if technical:
                    paragraph(routing_caption(frame.get("routing")), "Caption")
        else:
            paragraph("未附加人脸图片 现有缓存缺失或原素材已删除 导出不重新检测", "Caption")
        if large:
            doc.add_page_break()
        paragraph("完整逐帧记录" if record["upload"]["media_type"] == "video" else "图片检测明细", "Heading 1")
        paragraph("序号从 1 起计 时间单位为秒 分数按百分比显示 无人脸和未完成项不计为零", "Caption")
        rows = frame_rows(extra, technical)
        if len(rows) > 1:
            table(rows, [0.44, 0.61, 0.94, 0.69, 0.86, 0.64, 0.64, 0.64, 0.64, 1.06] if technical else [0.7, 1.1, 2.1, 1.25, 2.05])
        else:
            paragraph("此记录没有保存逐帧明细")
    for line in ending(data):
        paragraph(line, "Caption")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("智鉴分析报告  第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页")
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()
