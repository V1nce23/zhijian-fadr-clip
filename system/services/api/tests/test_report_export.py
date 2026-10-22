import csv
import io
import json
import uuid

import pytest

from app.core.database import SessionLocal
from app.models import Analysis, DetectorRun, Upload
from app.services import attention_gallery as attention
from app.services.routing_evidence import validated_routing
from app.services.storage import get_storage
from test_attention_gallery import payload
from test_evidence_gallery import report, files, publish_gallery, route
from test_live_visualization import running, prepare, gallery


def export_url(identity, format, technical=True):
    return f"/api/v1/analyses/{identity}/export?format={format}&include_technical={str(technical).lower()}"


def test_technical_details_are_optional_in_all_five_formats(client):
    from docx import Document
    from pypdf import PdfReader

    identity, run_id, _ = report(client)
    add_routing(run_id, live=True)
    before = client.get(f"/api/v1/analyses/{identity}").json()
    basic = client.get(f"/api/v1/analyses/{identity}/export?format=json").json()
    full = client.get(export_url(identity, "json")).json()
    assert basic["technical_details_included"] is False
    assert full["technical_details_included"] is True
    assert full["record"] == before
    assert basic["record_sha256"] == full["record_sha256"]
    assert basic["record"]["threshold_policy"] == before["threshold_policy"]
    assert basic["record"]["score"] == full["record"]["score"]
    assert "evidence" not in basic["record"]["detector_runs"][0]
    assert "model_version" not in basic["record"]["detector_runs"][0]
    assert "preprocessing_version" not in basic["record"]["detector_runs"][0]
    assert "routing_mean" not in basic["supplementary"][0]
    assert len(basic["supplementary"][0]["frames"]) == 3
    assert all("routing" not in frame for frame in basic["supplementary"][0]["frames"])
    assert full["supplementary"][0]["frames"][0]["routing"]["weights"] == [0.1, 0.2, 0.3, 0.4]
    for technical in (False, True):
        csv_body = client.get(export_url(identity, "csv", technical)).content.decode("utf-8-sig")
        rows = list(csv.DictReader(io.StringIO(csv_body)))
        assert ("模型版本" in rows[0]) is technical
        assert ("L24路由权重" in rows[0]) is technical
        assert rows[2]["帧分数（0–1）"] == ""
        assert rows[0]["帧分数（0–1）"] == "0.2"
        html = client.get(export_url(identity, "html", technical)).text
        assert ("模型版本" in html) is technical
        assert ("有效人脸帧平均路由权重" in html) is technical
        pdf = PdfReader(io.BytesIO(client.get(export_url(identity, "pdf", technical)).content))
        pdf_text = "\n".join(page.extract_text() for page in pdf.pages)
        assert ("模型版本" in pdf_text) is technical
        assert ("L24" in pdf_text) is technical
        assert "无人脸" in pdf_text
        word = Document(io.BytesIO(client.get(export_url(identity, "docx", technical)).content))
        assert ("模型版本" in "\n".join(p.text for p in word.paragraphs)) is technical
        assert len(word.tables[0].columns) == (10 if technical else 5)
        assert len(word.tables[0].rows) == 4
        assert word.tables[0].cell(3, 2).text == "无人脸，已跳过"
    assert client.get(f"/api/v1/analyses/{identity}").json() == before


def add_routing(run_id, *, live=False):
    with SessionLocal() as db:
        run = db.get(DetectorRun, uuid.UUID(run_id))
        evidence = {**run.evidence, "routing_layers": [6, 12, 18, 24], "routing_weights": [0.1, 0.2, 0.3, 0.4],
            "thresholds": {"authentic_max": 0.2, "manipulated_min": 0.2, "operator": ">"}}
        evidence["frames"] = [{**item, "routing_weights": [0.1, 0.2, 0.3, 0.4]} for item in evidence["frames"]]
        if live:
            evidence["visualization"] = {"version": "live-v1", "stage": "completed", "frames": [
                {**item, "status": "scored", "has_face": True, "has_source": False} for item in evidence["frames"]] + [
                {"frame_index": 20, "timestamp_seconds": 0.8, "status": "no_face", "score": None, "has_face": False, "has_source": False}]}
        run.evidence = evidence
        db.commit()


def test_exports_existing_record_without_queue_inference_or_report_mutation(client, tmp_path, monkeypatch):
    identity, run_id, _ = report(client)
    add_routing(run_id, live=True)
    publish_gallery(identity, run_id, files(tmp_path))
    before = client.get(f"/api/v1/analyses/{identity}").json()
    monkeypatch.setattr("app.detectors.fadr_clip.adapter.FADRClipDetector.load", lambda *args: pytest.fail("export must not infer"))
    monkeypatch.setattr("app.tasks.attention_gallery.prepare_attention_gallery.apply_async", lambda **kwargs: pytest.fail("export must not enqueue"))
    monkeypatch.setattr("app.tasks.evidence_gallery.prepare_evidence_gallery.apply_async", lambda **kwargs: pytest.fail("export must not extract"))
    for format in ("json", "csv", "html", "pdf", "docx"):
        response = client.get(export_url(identity, format))
        assert response.status_code == 200
        assert f'zhijian-report-{identity}.{format}' in response.headers["content-disposition"]
        assert response.headers["cache-control"] == "private, no-store"
    data = client.get(export_url(identity, "json")).json()
    assert data["record"] == before
    assert len(data["record_sha256"]) == 64
    frames = data["supplementary"][0]["frames"]
    assert [item["score"] for item in frames] == [0.2, 0.8, None]
    # Strict >: the exact threshold tie is authentic, never rounded first.
    assert [item["verdict"] for item in frames] == ["likely_authentic", "likely_manipulated", None]
    assert frames[0]["routing"]["weights"] == [0.1, 0.2, 0.3, 0.4]
    assert frames[2]["routing"] is None
    document = client.get(export_url(identity, "html")).text
    assert 'data:image/png;base64,' in document and "打印／另存为 PDF" in document
    assert 'http://storage.invalid' not in document
    assert client.get(f"/api/v1/analyses/{identity}").json() == before


def test_direct_pdf_embeds_chinese_font_and_word_has_editable_tables(client, tmp_path):
    from pypdf import PdfReader
    from docx import Document
    from zipfile import ZipFile

    identity, run_id, _ = report(client)
    add_routing(run_id, live=True)
    publish_gallery(identity, run_id, files(tmp_path))
    pdf = client.get(export_url(identity, "pdf"))
    assert pdf.headers["content-type"] == "application/pdf"
    assert pdf.content.startswith(b"%PDF-")
    reader = PdfReader(io.BytesIO(pdf.content))
    content = "\n".join(page.extract_text() for page in reader.pages)
    assert "智鉴伪造人脸分析报告" in content
    assert "无人脸" in content and "L24" in content and identity in content
    fonts = [font.get_object() for page in reader.pages for font in page["/Resources"]["/Font"].get_object().values()]
    assert any("/FontFile2" in font.get("/FontDescriptor", {}).get_object() for font in fonts if "/FontDescriptor" in font)
    word = client.get(export_url(identity, "docx"))
    assert word.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    document = Document(io.BytesIO(word.content))
    assert len(document.tables) == 1
    assert len(document.tables[0].rows) == 4
    assert document.tables[0].cell(3, 2).text == "无人脸，已跳过"
    assert len(document.inline_shapes) >= 2
    with ZipFile(io.BytesIO(word.content)) as archive:
        xml = archive.read("word/document.xml").decode()
        assert "w:tblHeader" in xml and "w:tblBorders" in xml
        assert 'w:hRule="exact"' not in xml
        assert b"w:pBdr" not in archive.read("word/styles.xml")


def test_csv_keeps_zero_scores_and_skipped_frames_distinct(client):
    identity, run_id, _ = report(client)
    add_routing(run_id, live=True)
    with SessionLocal() as db:
        run = db.get(DetectorRun, uuid.UUID(run_id))
        evidence = {**run.evidence}
        evidence["frames"] = [{**item, "score": 0.0} if item["frame_index"] == 0 else item for item in evidence["frames"]]
        run.evidence = evidence
        db.commit()
    response = client.get(export_url(identity, "csv"))
    assert response.content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert rows[0]["帧分数（0–1）"] == "0.0"
    assert rows[2]["帧分数（0–1）"] == "" and rows[2]["帧状态"] == "无人脸，已跳过"
    assert rows[0]["L24路由权重"] == "0.4" and rows[0]["画面序号（从1起）"] == "1"


def test_old_attention_json_routes_are_labeled_as_replay_and_never_replace_report_mean(client):
    identity, run_id, _ = report(client)
    attention.publish_attention(identity, run_id, index=0, payload=payload())
    storage = client.app.dependency_overrides[get_storage]()
    key = attention.manifest_key(identity, run_id)
    manifest = storage.read_json(key)
    # Simulate attention manifests written by v0.11 before routes were in manifest.
    manifest["frames"][0].pop("routing_weights")
    manifest["frames"][0].pop("routing_layers")
    storage.upload(key, io.BytesIO(json.dumps(manifest).encode()), "application/json")
    before = client.get(f"/api/v1/analyses/{identity}").json()
    data = client.get(export_url(identity, "json")).json()
    extra = data["supplementary"][0]
    assert extra["routing_mean"] is None
    assert extra["frames"][0]["routing"]["origin"] == "supplementary_replay"
    assert extra["frames"][1]["routing"] is None
    assert client.get(f"/api/v1/analyses/{identity}").json() == before


def test_live_scores_publish_routes_even_without_attention_assets(client):
    identity, run_id, reporter = running(client)
    prepare(reporter)
    reporter("scores", frames=[{"frame_index": 0, "score": 0.89, "routing_weights": [0.1, 0.2, 0.3, 0.4], "routing_layers": [6, 12, 18, 24]}])
    result = gallery(client, identity, run_id)
    assert result["frames"][0]["routing"]["weights"] == [0.1, 0.2, 0.3, 0.4]
    assert result["frames"][0]["heatmap_url"] is None and result["frames"][1]["routing"] is None


def test_escaped_html_and_csv_formula_injection_protection(client):
    identity, _, upload_id = report(client)
    with SessionLocal() as db:
        db.get(Upload, uuid.UUID(upload_id)).original_name = '=HYPERLINK("https://evil.invalid")<script>alert(1)</script>.mp4'
        db.get(Analysis, uuid.UUID(identity)).summary = '<img src=x onerror="alert(1)">'
        db.commit()
    document = client.get(export_url(identity, "html")).text
    assert "<script>alert(1)</script>" not in document and '<img src=x' not in document
    assert '&lt;script&gt;' in document and '&lt;img src=x' in document
    rows = list(csv.DictReader(io.StringIO(client.get(export_url(identity, "csv")).content.decode("utf-8-sig"))))
    assert rows[0]["文件名"].startswith("'=HYPERLINK")


@pytest.mark.parametrize("weights,layers", [([0.1, 0.2, 0.3, 0.5], [6, 12, 18, 24]), ([float("nan"), 0, 0, 1], [6, 12, 18, 24]), ([0.25] * 4, [1, 2, 3, 4]), ([True, 0, 0, 0], [6, 12, 18, 24])])
def test_invalid_routing_is_never_drawn(weights, layers):
    assert validated_routing(weights, layers) is None


def test_export_authentication_owner_deleted_source_and_missing_cache(client, monkeypatch):
    assert client.get(export_url(uuid.uuid4(), "html")).status_code == 401
    identity, run_id, upload_id = report(client)
    add_routing(run_id)
    assert client.get(export_url(identity, "exe")).status_code == 422
    assert client.delete(f"/api/v1/uploads/{upload_id}").status_code == 200
    # Report export is still useful after original media deletion.
    assert client.get(export_url(identity, "html")).status_code == 200
    assert client.get(export_url(identity, "pdf")).status_code == 200
    assert client.get(export_url(identity, "docx")).status_code == 200
    exported = client.get(export_url(identity, "json")).json()
    assert exported["supplementary"][0]["frames"][0]["routing"]["weights"] == [0.1, 0.2, 0.3, 0.4]
    client.post("/api/v1/users", json={"email": "export@example.com", "name": "Other", "password": "ExportPass123!", "role": "analyst"})
    client.post("/api/v1/auth/login", json={"email": "export@example.com", "password": "ExportPass123!"})
    assert client.get(export_url(identity, "json")).status_code == 404


def test_optional_storage_failure_does_not_prevent_exporting_database_record(client, monkeypatch):
    identity, run_id, _ = report(client)
    before = client.get(f"/api/v1/analyses/{identity}").json()
    storage = client.app.dependency_overrides[get_storage]()
    monkeypatch.setattr(storage, "exists", lambda *args: (_ for _ in ()).throw(OSError("offline cache")))
    response = client.get(export_url(identity, "html"))
    assert response.status_code == 200 and "未附加人脸图片" in response.text
    assert client.get(export_url(identity, "pdf")).status_code == 200
    assert client.get(export_url(identity, "docx")).status_code == 200
    assert client.get(f"/api/v1/analyses/{identity}").json() == before
